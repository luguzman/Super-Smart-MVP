#!/usr/bin/env python3
"""Importa el catálogo Excel a un modelo normalizado y auditable."""
from __future__ import annotations

import argparse
import json
import re
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


STOP_WORDS = {
    "de", "del", "la", "el", "los", "las", "y", "con", "sin", "para", "en",
    "por", "al", "un", "una", "unos", "unas", "aprox", "fresco", "fresca",
}


def as_number(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "."))
    except ValueError:
        return None


def as_text(value):
    return "" if value is None else str(value).strip()


def tokens(value: str) -> list[str]:
    value = value.lower()
    words = re.findall(r"[a-záéíóúñü0-9]+", value)
    return [word for word in words if word not in STOP_WORDS and len(word) > 1]


def read_sheet(workbook, sheet_name: str):
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    headers = [as_text(value) for value in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, default=Path("data"))
    args = parser.parse_args()

    args.output.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(args.source, read_only=True, data_only=True)
    products_sheet = read_sheet(workbook, "Productos")
    nutrition_sheet = read_sheet(workbook, "Nutricion")
    ingredients_sheet = read_sheet(workbook, "Ingredientes")

    nutrition = {as_text(row.get("product_id")): row for row in nutrition_sheet}
    ingredients = {as_text(row.get("product_id")): row for row in ingredients_sheet}
    products, observations, flags = [], [], []

    for row in products_sheet:
        product_id = as_text(row.get("product_id"))
        if not product_id:
            continue
        price = as_number(row.get("precio"))
        weight_g = as_number(row.get("presentacion_g"))
        portion_g = as_number(row.get("porcion_declarada_g"))
        servings = as_number(row.get("porciones_por_envase"))
        store = as_text(row.get("tienda"))
        product = {
            "id": product_id,
            "name": as_text(row.get("producto_nombre")),
            "brand": as_text(row.get("marca")),
            "store": store,
            "category": as_text(row.get("categoria")),
            "weight_g": weight_g,
            "ean": as_text(row.get("codigo_barras")),
            "notes": as_text(row.get("notas")),
            "source_url": as_text(row.get("url_fuente")),
            "name_tokens": tokens(as_text(row.get("producto_nombre"))),
            "nutrition": {key: value for key, value in nutrition.get(product_id, {}).items() if key and key != "product_id"},
            "ingredients": {key: value for key, value in ingredients.get(product_id, {}).items() if key and key != "product_id"},
        }
        product["unit_price_eur"] = round(price * 1000 / weight_g, 4) if price and weight_g else None
        products.append(product)

        if price and store:
            observations.append({
                "id": f"catalog-{product_id}",
                "product_id": product_id,
                "store": store,
                "price_eur": price,
                "observed_on": None,
                "promotion_eur": 0,
                "loyalty_required": False,
                "source": "catalogo_historico",
                "confidence": "media" if weight_g else "baja",
            })

        # En parte del fichero los dos campos contienen precio/precio unitario,
        # no gramos/porciones. Se conserva el original pero se impide usarlo.
        if portion_g and servings and (
            (price is not None and abs(portion_g - price) < 0.02)
            or (weight_g and portion_g * servings / weight_g > 1.35)
            or (weight_g and portion_g * servings / weight_g < 0.65)
        ):
            flags.append({
                "product_id": product_id,
                "field": "porcion_declarada_g/porciones_por_envase",
                "reason": "Valores semánticamente incompatibles con peso y/o precio; no se usan en el optimizador.",
            })

    quality = {
        "generated_on": date.today().isoformat(),
        "products": len(products),
        "observations": len(observations),
        "missing_price": sum(1 for product in products if not any(obs["product_id"] == product["id"] for obs in observations)),
        "missing_ean": sum(1 for product in products if not product["ean"]),
        "missing_source_url": sum(1 for product in products if not product["source_url"]),
        "semantic_flags": flags,
        "disclaimer": "Los precios importados proceden del catálogo histórico y requieren actualización antes de una compra real.",
    }
    (args.output / "catalog.json").write_text(json.dumps(products, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.output / "price_observations.json").write_text(json.dumps(observations, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.output / "quality_report.json").write_text(json.dumps(quality, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Importados {len(products)} productos y {len(observations)} observaciones.")


if __name__ == "__main__":
    main()
